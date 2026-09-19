from copy import copy
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook


QC_SHEET_NAME = "Lims-品管數據"
ANALYSIS_SHEET_NAME = "Lims-分析數據"

DATA_START_ROW = 5


def _copy_row_style(
    worksheet,
    source_row,
    target_row,
    max_column
):

    for column in range(
        1,
        max_column + 1
    ):

        source_cell = worksheet.cell(
            row=source_row,
            column=column
        )

        target_cell = worksheet.cell(
            row=target_row,
            column=column
        )

        if source_cell.has_style:

            target_cell._style = copy(
                source_cell._style
            )

        if source_cell.number_format:

            target_cell.number_format = (
                source_cell.number_format
            )

        if source_cell.font:

            target_cell.font = copy(
                source_cell.font
            )

        if source_cell.fill:

            target_cell.fill = copy(
                source_cell.fill
            )

        if source_cell.border:

            target_cell.border = copy(
                source_cell.border
            )

        if source_cell.alignment:

            target_cell.alignment = copy(
                source_cell.alignment
            )

        if source_cell.protection:

            target_cell.protection = copy(
                source_cell.protection
            )


def _clear_data_rows(
    worksheet,
    start_row,
    max_column
):

    for row_number in range(
        start_row,
        worksheet.max_row + 1
    ):

        for column in range(
            1,
            max_column + 1
        ):

            worksheet.cell(
                row=row_number,
                column=column
            ).value = None


def _parse_analysis_date(
    analysis_date
):

    if analysis_date is None:
        return None

    if isinstance(
        analysis_date,
        datetime
    ):
        return analysis_date

    date_text = str(
        analysis_date
    ).strip()

    if not date_text:
        return None

    supported_formats = [
        "%Y-%m-%d",
        "%Y/%m/%d"
    ]

    for date_format in supported_formats:

        try:

            return datetime.strptime(
                date_text,
                date_format
            )

        except ValueError:
            continue

    return date_text


def build_lims_excel(
    template_path,
    export_data,
    analyst_id="",
    analyst_name="",
    analysis_method=""
):

    workbook = load_workbook(
        template_path,
        keep_links=True
    )

    if QC_SHEET_NAME not in workbook.sheetnames:

        raise ValueError(
            "LIMS 範本缺少工作表："
            + QC_SHEET_NAME
        )

    if ANALYSIS_SHEET_NAME not in workbook.sheetnames:

        raise ValueError(
            "LIMS 範本缺少工作表："
            + ANALYSIS_SHEET_NAME
        )

    qc_sheet = workbook[
        QC_SHEET_NAME
    ]

    analysis_sheet = workbook[
        ANALYSIS_SHEET_NAME
    ]

    # ========================================
    # 取得分析日期
    # 優先使用第一筆品管資料
    # ========================================

    analysis_date = None

    if export_data.qc_rows:

        analysis_date = (
            export_data
            .qc_rows[0]
            .analysis_date
        )

    analysis_date_value = (
        _parse_analysis_date(
            analysis_date
        )
    )

    # ========================================
    # 表頭共同資料
    # ========================================

    qc_sheet["B2"] = (
        analysis_date_value
    )

    analysis_sheet["B2"] = (
        analysis_date_value
    )

    qc_sheet["B3"] = analyst_id
    qc_sheet["C3"] = analyst_name

    analysis_sheet["B3"] = analyst_id
    analysis_sheet["C3"] = analyst_name

    # ========================================
    # 分析方法
    # 僅 Lims-品管數據使用
    # ========================================

    qc_sheet["C1"] = analysis_method

    # 日期保持 Excel 日期格式
    if isinstance(
        analysis_date_value,
        datetime
    ):

        qc_sheet["B2"].number_format = (
            "yyyy/mm/dd"
        )

        analysis_sheet[
            "B2"
        ].number_format = (
            "yyyy/mm/dd"
        )

    # ========================================
    # 清除範本既有測試資料
    # 不刪除列，避免破壞範本結構
    # ========================================

    _clear_data_rows(
        qc_sheet,
        DATA_START_ROW,
        11
    )

    _clear_data_rows(
        analysis_sheet,
        DATA_START_ROW,
        6
    )

    # ========================================
    # 寫入 Lims-品管數據
    #
    # A 項目代碼
    # B 設備序號
    # C 最高吸光度
    # D 重複分析
    # E 查核分析
    # F 添加分析
    # G 檢量線確認
    # H 空白試驗
    # I 備註
    # J 批次流水號
    # K 檢量線係數
    # ========================================

    for index, row in enumerate(
        export_data.qc_rows,
        start=DATA_START_ROW
    ):

        if index > DATA_START_ROW:

            _copy_row_style(
                qc_sheet,
                DATA_START_ROW,
                index,
                11
            )

        qc_sheet.cell(
            row=index,
            column=1,
            value=row.exam_no
        )

        qc_sheet.cell(
            row=index,
            column=2,
            value=row.instrument_id
        )

        qc_sheet.cell(
            row=index,
            column=3,
            value=row.max_signal
        )

        duplicate_cell = qc_sheet.cell(
            row=index,
            column=4
        )

        if row.duplicate_result is not None:
            duplicate_cell.value = round(
            row.duplicate_result,
            1
            )
            duplicate_cell.number_format = "0.0"


        qc_cell = qc_sheet.cell(
            row=index,
            column=5
        )

        if row.qc_result is not None:
           qc_cell.value = round(
           row.qc_result,
           1
           )
           qc_cell.number_format = "0.0"


        spike_cell = qc_sheet.cell(
            row=index,
            column=6
           )

        if row.spike_result is not None:
            spike_cell.value = round(
              row.spike_result,
              1
            )
            spike_cell.number_format = "0.0"


        calibration_check_cell = qc_sheet.cell(
             row=index,
             column=7
            )

        if row.calibration_check is not None:
             calibration_check_cell.value = round(
             row.calibration_check,
              1
             )
             calibration_check_cell.number_format = "0.0"
       
        qc_sheet.cell(
            row=index,
            column=8,
            value=row.blank_result
        )

        qc_sheet.cell(
            row=index,
            column=9,
            value=row.remark
        )

        qc_sheet.cell(
            row=index,
            column=10,
            value=row.batch_no
        )

        correlation_cell = qc_sheet.cell(
            row=index,
            column=11
        )

        if row.correlation_coefficient is not None:

           correlation_cell.value = round(
           row.correlation_coefficient,
           4
         )

           correlation_cell.number_format = (
           "0.0000"
         )

    # ========================================
    # 寫入 Lims-分析數據
    #
    # A 樣品編號
    # B 項目代碼
    # C 檢驗結果
    # D 實際分析項目
    # E 批次流水號
    # F 備註
    # ========================================

    for index, row in enumerate(
        export_data.analysis_rows,
        start=DATA_START_ROW
    ):

        if index > DATA_START_ROW:

            _copy_row_style(
                analysis_sheet,
                DATA_START_ROW,
                index,
                6
            )

        analysis_sheet.cell(
            row=index,
            column=1,
            value=row.sample_id
        )

        analysis_sheet.cell(
            row=index,
            column=2,
            value=row.exam_no
        )

        analysis_sheet.cell(
            row=index,
            column=3,
            value=row.result
        )

        analysis_sheet.cell(
            row=index,
            column=4,
            value=row.actual_item
        )

        analysis_sheet.cell(
            row=index,
            column=5,
            value=row.batch_no
        )

        analysis_sheet.cell(
            row=index,
            column=6,
            value=row.remark
        )

    # ========================================
    # 輸出至記憶體
    # Flask send_file 可直接使用
    # ========================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    return output